# -*- coding: utf-8 -*-
"""
Learning Scout v2 - 报告生成器模块
生成格式化的学习报告（Markdown + DOCX）
"""
import os
import json
import sqlite3
from datetime import datetime, timedelta
from typing import List, Dict, Any, Optional

DB_PATH = os.path.join(os.path.dirname(__file__), "dapper_memory.db")


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _resolve_db():
    """兼容本地和VPS的数据库路径"""
    path = os.environ.get("DB_PATH", "")
    if path and os.path.exists(path):
        return path
    # VPS 路径
    vps_path = "/opt/Dapper_Coding_Agent/dapper_memory.db"
    if os.path.exists(vps_path):
        return vps_path
    return DB_PATH


def get_user_stats(user_id: str, days: int = 7) -> Dict[str, Any]:
    """获取用户最近N天的学习统计数据"""
    db = _resolve_db()
    cutoff = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:%M:%S")

    stats = {
        "user_id": user_id,
        "period_days": days,
        "lessons_generated": 0,
        "quizzes_taken": 0,
        "quiz_correct_rate": 0.0,
        "avg_mastery_score": 0.0,
        "topics_learned": [],
        "weak_topics": [],
        "strong_topics": [],
    }

    try:
        conn = sqlite3.connect(db)
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()

        # 课程生成数量
        cur.execute(
            "SELECT COUNT(*) FROM learning_history WHERE user_id=? AND created_at>=?",
            (user_id, cutoff)
        )
        stats["lessons_generated"] = cur.fetchone()[0]

        # Quiz 答题情况
        cur.execute(
            """SELECT COUNT(*), SUM(CASE WHEN is_correct=1 THEN 1 ELSE 0 END)
               FROM quiz_results WHERE user_id=? AND answered_at>=?""",
            (user_id, cutoff)
        )
        row = cur.fetchone()
        total = row[0] or 0
        correct = row[1] or 0
        stats["quizzes_taken"] = total
        stats["quiz_correct_rate"] = round(correct / total * 100, 1) if total > 0 else 0.0

        # Mastery score 统计
        cur.execute(
            """SELECT topic, mastery_score FROM mastery_records
               WHERE user_id=? ORDER BY updated_at DESC LIMIT 20""",
            (user_id,)
        )
        scores = cur.fetchall()
        if scores:
            score_vals = [s["mastery_score"] for s in scores]
            stats["avg_mastery_score"] = round(sum(score_vals) / len(score_vals), 2)
            stats["topics_learned"] = [s["topic"] for s in scores if s["mastery_score"] >= 0.5]
            stats["weak_topics"] = [s["topic"] for s in scores if s["mastery_score"] < 0.4]
            stats["strong_topics"] = [s["topic"] for s in scores if s["mastery_score"] >= 0.8]

        conn.close()
    except Exception as e:
        stats["error"] = str(e)

    return stats


def generate_markdown_report(user_id: str, stats: Dict[str, Any]) -> str:
    """生成 Markdown 格式的学习报告"""
    lines = [
        f"# 📚 学习周报",
        f"**用户**: {user_id}",
        f"**生成时间**: {_now()}",
        f"**统计周期**: 近 {stats['period_days']} 天",
        "",
        "---",
        "",
        "## 📊 学习数据一览",
        "",
        f"| 指标 | 数值 |",
        f"|------|------|",
        f"| 课程生成数 | {stats['lessons_generated']} |",
        f"| 答题数 | {stats['quizzes_taken']} |",
        f"| 正确率 | {stats['quiz_correct_rate']}% |",
        f"| 平均掌握度 | {stats['avg_mastery_score']} |",
        "",
    ]

    if stats.get("strong_topics"):
        lines += [
            "## 🟢 已掌握知识点",
            "",
            ", ".join(f"`{t}`" for t in stats["strong_topics"]),
            "",
        ]

    if stats.get("weak_topics"):
        lines += [
            "## 🔴 待加强知识点",
            "",
            ", ".join(f"`{t}`" for t in stats["weak_topics"]),
            "",
        ]

    lines += [
        "## 💡 学习建议",
        "",
        _generate_advice(stats),
        "",
        "---",
        "*由 Learning Scout 自动生成*",
    ]

    return "\n".join(lines)


def _generate_advice(stats: Dict[str, Any]) -> str:
    """基于统计数据生成个性化建议"""
    advices = []

    if stats["quiz_correct_rate"] < 60:
        advices.append("本周正确率偏低，建议复习之前课程的例题部分")
    elif stats["quiz_correct_rate"] >= 80:
        advices.append("正确率不错，可以挑战更难的题目了")

    if len(stats.get("weak_topics", [])) > 3:
        advices.append(f"你有 {len(stats['weak_topics'])} 个薄弱知识点，建议每天复习一个")
    elif len(stats.get("weak_topics", [])) == 0 and stats["quizzes_taken"] > 0:
        advices.append("所有已学知识点都掌握良好，继续探索新主题吧")

    if stats["lessons_generated"] == 0:
        advices.append("本周还没有生成新课程，试试发送「给我讲讲 Python 异步编程」")
    elif stats["lessons_generated"] >= 5:
        advices.append("学习进度很快，注意复习巩固哦")

    return advices[0] if advices else "保持学习节奏，持续进步！"


def generate_docx_report(user_id: str, stats: Dict[str, Any], output_path: str = None) -> str:
    """生成 DOCX 格式的学习报告"""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return ""

    if output_path is None:
        reports_dir = os.path.join(os.path.dirname(__file__), "reports")
        os.makedirs(reports_dir, exist_ok=True)
        output_path = os.path.join(reports_dir, f"report_{user_id}_{datetime.now().strftime('%Y%m%d')}.docx")

    doc = Document()

    # 标题
    title = doc.add_heading("📚 Learning Scout 学习周报", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 基本信息
    doc.add_paragraph(f"用户: {user_id}")
    doc.add_paragraph(f"生成时间: {_now()}")
    doc.add_paragraph(f"统计周期: 近 {stats['period_days']} 天")

    doc.add_heading("📊 学习数据一览", level=2)

    table = doc.add_table(rows=5, cols=2)
    table.style = "Light Grid Accent 1"
    data = [
        ("指标", "数值"),
        ("课程生成数", str(stats["lessons_generated"])),
        ("答题数", str(stats["quizzes_taken"])),
        ("正确率", f"{stats['quiz_correct_rate']}%"),
        ("平均掌握度", str(stats["avg_mastery_score"])),
    ]
    for i, (k, v) in enumerate(data):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v

    if stats.get("strong_topics"):
        doc.add_heading("🟢 已掌握知识点", level=2)
        doc.add_paragraph(", ".join(stats["strong_topics"]))

    if stats.get("weak_topics"):
        doc.add_heading("🔴 待加强知识点", level=2)
        doc.add_paragraph(", ".join(stats["weak_topics"]))

    doc.add_heading("💡 学习建议", level=2)
    doc.add_paragraph(_generate_advice(stats))

    doc.add_paragraph("---")
    p = doc.add_paragraph("*由 Learning Scout 自动生成*")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    # 测试
    user_id = "test_user"
    stats = get_user_stats(user_id)
    print("Stats:", json.dumps(stats, ensure_ascii=False, indent=2))

    md = generate_markdown_report(user_id, stats)
    print("\nMarkdown Report:")
    print(md)
